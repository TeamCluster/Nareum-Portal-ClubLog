"""지난 활동이력(엑셀) 일회성 승계 마이그레이션.

목적
----
시스템 가동 이전의 과거 활동이력이 엑셀(이 시스템의 다운로드 포맷과 동일)로
남아있다. 이를 각 기관 DB(db/<slug>.sqlite3)의 logs 테이블로 들여오되:

  * 가장 과거 기록이 id = 1 이 되도록 (기록일시 오름차순)
  * 기존 라이브 기록과 합쳐 전체를 created_at 순으로 재정렬 후 id 1..N 재부여
  * 엑셀의 '기록일시'(created_at), '활동 일자'(activity_date) 원본을 그대로 보존

백엔드 코드는 건드리지 않는다. 이 스크립트가 sqlite 파일을 직접 재작성한다.

입력 규칙
---------
    import/<slug>/*.xlsx

  - 폴더명이 기관 slug. (예: import/nareum/2026_상반기.xlsx)
  - 한 기관에 엑셀이 여러 개면 폴더에 모두 넣으면 합쳐서 처리.
  - 엑셀 컬럼은 다운로드 포맷(EXCEL_HEADERS)과 동일한 한글 헤더를 기대.

사용법
------
    # 1) 미리보기 (아무것도 안 바꿈 — 파싱/정렬/중복 결과만 출력)
    python migrate_import_history.py

    # 2) 실제 반영 (실행 전 db/_backup_<시각>/ 에 자동 백업)
    python migrate_import_history.py --commit

    # 특정 기관만:  python migrate_import_history.py --slug nareum --commit
"""
import argparse
import os
import shutil
import sqlite3
import sys
from datetime import datetime, date, time

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("openpyxl 가 필요합니다.  pip install openpyxl==3.1.5")

# Windows 콘솔(cp949)에서 한글 깨짐 방지
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

HERE = os.path.dirname(os.path.abspath(__file__))
DB_FOLDER = os.path.join(HERE, "db")
IMPORT_ROOT = os.path.join(HERE, "import")

# 엑셀 한글 헤더 -> logs 컬럼.  log_service.EXCEL_HEADERS 의 역방향.
HEADER_TO_COL = {
    "동아리 분야": "category",
    "동아리 명":   "club_name",
    "활동 일자":   "activity_date",
    "시작 시간":   "start_time",
    "종료 시간":   "end_time",
    "참가자":      "participants",
    "활동 내용":   "content",
    "작성자":      "author",
    "총 인원수":   "total_count",
    "초등 남":     "elem_m",
    "초등 여":     "elem_f",
    "중등 남":     "mid_m",
    "중등 여":     "mid_f",
    "고등 남":     "high_m",
    "고등 여":     "high_f",
    "후기 남":     "univ_m",
    "후기 여":     "univ_f",
    "기록일시":    "created_at",
}

COUNT_COLS = ["elem_m", "elem_f", "mid_m", "mid_f", "high_m", "high_f", "univ_m", "univ_f"]

# INSERT 시 사용할 컬럼 순서 (id 제외 — 재부여)
INSERT_COLS = [
    "category", "club_name", "activity_date", "start_time", "end_time",
    "participants", "content", "author", "total_count",
    *COUNT_COLS, "created_at",
]


# ----------------------------------------------------------------------
# 값 정규화 — 엑셀 셀은 문자열/숫자/날짜/시각 무엇이든 올 수 있음
# ----------------------------------------------------------------------
def norm_date(v):
    """활동 일자 -> 'YYYY-MM-DD'."""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    # '2026-02-01', '2026-02-01 00:00:00', '2026/2/1' 등 폭넓게 수용
    s = s.replace("/", "-").split(" ")[0]
    y, m, d = s.split("-")
    return f"{int(y):04d}-{int(m):02d}-{int(d):02d}"


def norm_time(v):
    """시작/종료 시간 -> 'HH:MM'.  '12:0', '13:50', time 객체 등 수용."""
    if isinstance(v, datetime):
        return v.strftime("%H:%M")
    if isinstance(v, time):
        return v.strftime("%H:%M")
    s = str(v).strip()
    if not s:
        return ""
    parts = s.split(":")
    h = int(parts[0])
    mi = int(parts[1]) if len(parts) > 1 and parts[1] != "" else 0
    return f"{h:02d}:{mi:02d}"


def norm_created_at(v):
    """기록일시 -> 'YYYY-MM-DD HH:MM:SS' (원본 보존, 형식만 정규화)."""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d 00:00:00")
    s = str(v).strip()
    s = s.replace("/", "-").replace("T", " ")
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            continue
    raise ValueError(f"기록일시 형식을 해석할 수 없음: {v!r}")


def norm_int(v):
    if v is None or v == "":
        return 0
    return int(float(v))


def norm_text(v):
    return "" if v is None else str(v).strip()


# ----------------------------------------------------------------------
# 엑셀 -> 레코드 dict 리스트
# ----------------------------------------------------------------------
def read_excel(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)

    try:
        header = next(rows)
    except StopIteration:
        return []

    # 헤더 -> 열 인덱스
    col_index = {}
    for idx, name in enumerate(header):
        key = str(name).strip() if name is not None else ""
        if key in HEADER_TO_COL:
            col_index[HEADER_TO_COL[key]] = idx

    missing = [h for h, c in HEADER_TO_COL.items()
               if c not in col_index and c not in ("participants", "content", "author", "category")]
    if "club_name" not in col_index or "created_at" not in col_index:
        raise ValueError(
            f"{os.path.basename(path)}: 필수 헤더(동아리 명/기록일시)를 찾지 못함. "
            f"발견된 헤더: {[str(h) for h in header]}"
        )
    if missing:
        print(f"    ⚠ {os.path.basename(path)}: 일부 헤더 없음 {missing} (0/빈값 처리)")

    out = []
    derived = 0
    for r in rows:
        if r is None or all(c is None or str(c).strip() == "" for c in r):
            continue  # 빈 줄 스킵

        def cell(colname):
            i = col_index.get(colname)
            return r[i] if i is not None and i < len(r) else None

        rec = {
            "category":      norm_text(cell("category")),
            "club_name":     norm_text(cell("club_name")),
            "activity_date": norm_date(cell("activity_date")),
            "start_time":    norm_time(cell("start_time")),
            "end_time":      norm_time(cell("end_time")),
            "participants":  norm_text(cell("participants")),
            "content":       norm_text(cell("content")),
            "author":        norm_text(cell("author")),
        }

        # 기록일시 — 비어있으면(시스템 초기 데이터) 활동일자+종료시간으로 합성.
        # 시간순 정렬·NOT NULL 제약을 위해 필요. 합성 건수는 리포트한다.
        raw_created = cell("created_at")
        if raw_created is None or str(raw_created).strip() == "":
            t = rec["end_time"] or rec["start_time"] or "00:00"
            rec["created_at"] = f"{rec['activity_date']} {t}:00"
            rec["_derived_created"] = True
            derived += 1
        else:
            rec["created_at"] = norm_created_at(raw_created)
            rec["_derived_created"] = False

        for c in COUNT_COLS:
            rec[c] = norm_int(cell(c))

        computed = sum(rec[c] for c in COUNT_COLS)
        excel_total = norm_int(cell("total_count"))
        rec["total_count"] = computed  # 8개 합을 신뢰 (DB 제약: total_count NOT NULL)
        if excel_total and excel_total != computed:
            print(f"    ⚠ 총인원 불일치 [{rec['club_name']} {rec['activity_date']}]: "
                  f"엑셀={excel_total} 계산={computed} → 계산값 사용")
        out.append(rec)

    wb.close()
    if derived:
        print(f"    ⚠ {os.path.basename(path)}: 기록일시 없는 {derived}건 → "
              f"활동일자+종료시간으로 합성")
    return out


# ----------------------------------------------------------------------
# DB 입출력
# ----------------------------------------------------------------------
def fetch_existing(db_path):
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    rows = conn.execute("SELECT * FROM logs").fetchall()
    conn.close()
    return [{k: row[k] for k in row.keys() if k != "id"} for row in rows]


def dedup_key(rec):
    return (rec["created_at"], rec["club_name"], rec["activity_date"], rec["content"])


def sort_key(rec):
    # 기록일시 우선, 동률이면 활동일자·시작시간으로 안정 정렬
    return (rec["created_at"], rec["activity_date"], rec["start_time"])


def rewrite_logs(db_path, ordered):
    conn = sqlite3.connect(db_path)
    try:
        conn.execute("BEGIN")
        conn.execute("DELETE FROM logs")
        # AUTOINCREMENT 카운터 리셋 → 새 INSERT 가 id 1 부터
        conn.execute("DELETE FROM sqlite_sequence WHERE name = 'logs'")
        placeholders = ", ".join("?" for _ in INSERT_COLS)
        sql = f"INSERT INTO logs ({', '.join(INSERT_COLS)}) VALUES ({placeholders})"
        conn.executemany(sql, [[rec[c] for c in INSERT_COLS] for rec in ordered])
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


# ----------------------------------------------------------------------
# 기관 1곳 처리
# ----------------------------------------------------------------------
def process_slug(slug, commit, backup_dir):
    folder = os.path.join(IMPORT_ROOT, slug)
    db_path = os.path.join(DB_FOLDER, f"{slug}.sqlite3")
    print(f"\n=== {slug} ===")
    if not os.path.exists(db_path):
        print(f"  ✗ DB 없음: {db_path} — 건너뜀")
        return

    xlsx = sorted(f for f in os.listdir(folder)
                  if f.lower().endswith(".xlsx") and not f.startswith("~$"))
    if not xlsx:
        print(f"  (import/{slug}/ 에 .xlsx 없음 — 건너뜀)")
        return

    # 1) 엑셀 과거 기록 읽기
    past = []
    for f in xlsx:
        recs = read_excel(os.path.join(folder, f))
        print(f"  · {f}: {len(recs)}건")
        past.extend(recs)

    # 2) 기존 라이브 기록
    existing = fetch_existing(db_path)
    print(f"  · 기존 DB 라이브 기록: {len(existing)}건")

    # 3) 중복 제거: 기존에 이미 있는 과거행은 스킵 + 엑셀 내부 중복도 제거
    seen = {dedup_key(r) for r in existing}
    fresh, dup = [], 0
    for r in past:
        k = dedup_key(r)
        if k in seen:
            dup += 1
            continue
        seen.add(k)
        fresh.append(r)
    if dup:
        print(f"  · 중복 스킵: {dup}건")

    # 4) 병합 + 정렬
    merged = existing + fresh
    merged.sort(key=sort_key)

    print(f"  → 결과: 과거 {len(fresh)} + 기존 {len(existing)} = 총 {len(merged)}건, "
          f"id 1..{len(merged)} 재부여")
    if merged:
        first, last = merged[0], merged[-1]
        print(f"      id 1   : {first['created_at']}  [{first['club_name']}] {first['activity_date']}")
        print(f"      id {len(merged):<4}: {last['created_at']}  [{last['club_name']}] {last['activity_date']}")
    # 미리보기: 앞 3건
    for i, r in enumerate(merged[:3], 1):
        print(f"      예시 id {i}: {r['created_at']} | {r['category']}/{r['club_name']} | "
              f"{r['activity_date']} {r['start_time']}~{r['end_time']} | 총 {r['total_count']}명")

    if not commit:
        print("  (dry-run — 반영 안 함. 실제 적용하려면 --commit)")
        return

    # 5) 백업 후 재작성
    os.makedirs(backup_dir, exist_ok=True)
    shutil.copy2(db_path, os.path.join(backup_dir, f"{slug}.sqlite3"))
    rewrite_logs(db_path, merged)
    print(f"  ✓ 반영 완료. 백업: {os.path.relpath(os.path.join(backup_dir, slug + '.sqlite3'), HERE)}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--commit", action="store_true", help="실제 DB 반영 (없으면 dry-run)")
    ap.add_argument("--slug", help="특정 기관만 처리")
    args = ap.parse_args()

    if not os.path.isdir(IMPORT_ROOT):
        sys.exit(f"import 폴더가 없습니다: {IMPORT_ROOT}\n"
                 f"  import/<slug>/*.xlsx 형태로 엑셀을 넣어주세요.")

    if args.slug:
        slugs = [args.slug]
    else:
        slugs = sorted(d for d in os.listdir(IMPORT_ROOT)
                       if os.path.isdir(os.path.join(IMPORT_ROOT, d)))
    if not slugs:
        sys.exit("import/ 아래에 기관 폴더가 없습니다.")

    backup_dir = os.path.join(DB_FOLDER, "_backup_" + datetime.now().strftime("%Y%m%d_%H%M%S"))
    mode = "COMMIT" if args.commit else "DRY-RUN"
    print(f"[{mode}] 대상 기관: {', '.join(slugs)}")

    for slug in slugs:
        process_slug(slug, args.commit, backup_dir)

    print("\n완료.")


if __name__ == "__main__":
    main()
